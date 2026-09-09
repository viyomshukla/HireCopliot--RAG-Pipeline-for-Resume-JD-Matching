"""
Synthetic resume generator for the Recruiter / Hiring Copilot RAG project.

WHY THIS EXISTS (read before changing anything)
-----------------------------------------------
This script is not just "make some test files". It is the ground truth for
every stage downstream, so it has four jobs:

1. GROUND TRUTH. Every resume is generated FROM a structured record, and that
   record is written to candidates.json. Later, the extraction stage uses an
   LLM to recover skills/years/education from the rendered document. Without
   the original record you have nothing to score that against. This file is
   your free, exact evaluation set.

2. LEXICAL DIVERSITY. If every Software Engineer draws bullets from the same
   pool of 8 sentences, then every candidate's chunks embed to nearly the same
   vector, BM25 term frequencies go uniform, and retrieval quality becomes
   unmeasurable -- you cannot rank documents that are copies of each other.
   So responsibilities are TEMPLATED with randomised metrics, tools and
   domains, and skills appear under varied surface forms.

3. VOCABULARY DRIFT. The project's headline claim is "semantic search finds
   skills phrased differently than the JD says them". That claim is untestable
   if the resume and the JD both say "PyTorch". So each skill has a canonical
   name (stored in ground truth, used by JDs and hard filters) and several
   surface phrasings (what actually appears in the document). Dense retrieval
   has to bridge that gap; BM25 alone cannot.

4. FORMAT STRESS. Real resumes vary. Section headings use different aliases,
   sections get reordered or omitted, unknown sections appear, dates come in
   four formats, job titles sometimes share a line with their dates and
   sometimes do not, and DOCX headings are sometimes real Heading styles and
   sometimes just bold uppercase paragraphs. All of this is randomised per
   resume so the chunker is tested rather than flattered.

FAIRNESS NOTE
-------------
Proxy attributes (name locale, university prestige tier, career gaps) are
recorded ONLY in candidates.json, never stated in the document text. That is
exactly what a disparate-impact audit needs: a hidden attribute the ranker
cannot see directly but might correlate with anyway.

PRODUCTION vs LEARNING SHORTCUT
-------------------------------
Real systems evaluate on held-out real resumes with human relevance labels.
Synthetic data with perfect ground truth is a learning-stage substitute: it
lets you measure extraction accuracy exactly, but it will overstate parser
robustness, because a generator can only produce the failure modes you thought
to program. Treat green metrics here as a floor, not a ceiling.
"""

from __future__ import annotations

import argparse
import json
import random
from pathlib import Path

from faker import Faker

from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


# ============================================================
# OUTPUT PATHS
# ============================================================

BASE = Path(__file__).resolve().parents[1]
OUTPUT_DIR = BASE / "data" / "sample_resumes"
GROUND_TRUTH_FILE = BASE / "data" / "ground_truth" / "candidates.json"

CURRENT_YEAR = 2026


# ============================================================
# FONT HANDLING
# ============================================================
# ReportLab's built-in Helvetica does not embed a ToUnicode map for non-ASCII
# glyphs. Measured behaviour: "\u2022" extracts as "(cid:127)" and "\u25aa",
# "\u25e6", "\u2023" all extract as the LETTER "n" -- silent corruption with no
# error. Embedding a TrueType font makes every glyph round-trip correctly.
#
# If no TTF is found we fall back to Helvetica AND to ASCII-only bullets, so
# the data stays clean rather than becoming quietly wrong.

FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "C:/Windows/Fonts/calibri.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
]

FONT_CANDIDATES_BOLD = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    "C:/Windows/Fonts/arialbd.ttf",
    "C:/Windows/Fonts/calibrib.ttf",
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    "/Library/Fonts/Arial Bold.ttf",
]

UNICODE_BULLETS = ["\u2022", "\u25aa", "\u25e6", "\u2023", "-", "*"]
ASCII_BULLETS = ["-", "*"]


def register_fonts() -> tuple[str, str, list[str]]:
    """Returns (regular_font, bold_font, safe_bullet_set)."""
    regular = next((p for p in FONT_CANDIDATES if Path(p).exists()), None)
    bold = next((p for p in FONT_CANDIDATES_BOLD if Path(p).exists()), None)

    if regular and bold:
        pdfmetrics.registerFont(TTFont("ResumeFont", regular))
        pdfmetrics.registerFont(TTFont("ResumeFont-Bold", bold))
        pdfmetrics.registerFontFamily(
            "ResumeFont", normal="ResumeFont", bold="ResumeFont-Bold"
        )
        return "ResumeFont", "ResumeFont-Bold", UNICODE_BULLETS

    print(
        "WARNING: no embeddable TTF found. Falling back to Helvetica with "
        "ASCII-only bullets. Unicode bullets would silently corrupt on "
        "extraction (\u25aa becomes the letter 'n')."
    )
    return "Helvetica", "Helvetica-Bold", ASCII_BULLETS


# ============================================================
# ROLE PROFILES
# ============================================================
# Each responsibility is a TEMPLATE with {slots}. Filling the slots with random
# values is what stops 50 resumes from being 50 copies of the same 8 sentences.

CAREER_PROFILES: dict[str, dict] = {
    "Software Engineer": {
        "titles": [
            "Software Engineer",
            "Software Development Engineer",
            "Backend Software Engineer",
            "Senior Software Engineer",
        ],
        "skills": [
            "Python", "Java", "SQL", "Git", "REST APIs", "Docker",
            "Kubernetes", "Redis", "Microservices", "Unit Testing",
        ],
        "companies": [
            "Google", "Microsoft", "Amazon", "Infosys", "TCS", "Wipro",
            "Accenture", "Zoho", "Freshworks", "Cognizant",
        ],
        "responsibilities": [
            "Built and maintained {count} {lang} microservices handling {volume} requests per day for the {domain} platform.",
            "Reduced p95 API latency from {latency_before}ms to {latency_after}ms by adding {cache} caching and rewriting the hottest query path.",
            "Migrated the {domain} service from a monolith to {count} independently deployable services, cutting release time by {pct}%.",
            "Designed the {domain} data model in {db} and wrote the migration scripts for a {volume}-row backfill.",
            "Raised unit test coverage on the {domain} module from {pct_low}% to {pct}% and wired it into the CI gate.",
            "Containerised {count} legacy services with Docker and moved deployment to {orchestrator}.",
            "Reviewed roughly {count} pull requests per month and mentored {team_size} junior engineers.",
            "Debugged and fixed a production incident affecting {volume} users, then wrote the postmortem and added alerting.",
            "Introduced structured logging and tracing across the {domain} stack, cutting mean time to diagnosis by {pct}%.",
            "Implemented {auth} authentication and role-based access control for the internal {domain} tooling.",
        ],
        "projects": [
            ("Distributed URL Shortener", "A {lang} service with {cache}-backed counters, handling {volume} redirects per day."),
            ("E-commerce Order Backend", "REST APIs over {db} for cart, checkout and {domain}, deployed with Docker."),
            ("Realtime Chat Server", "WebSocket server in {lang} supporting {volume} concurrent connections."),
            ("CI Pipeline Automation", "Automated build and test pipeline that cut feedback time by {pct}%."),
        ],
    },
    "Data Analyst": {
        "titles": [
            "Data Analyst", "Business Data Analyst",
            "Senior Data Analyst", "Analytics Specialist",
        ],
        "skills": [
            "SQL", "Python", "Excel", "Power BI", "Tableau", "Statistics",
            "Data Visualization", "ETL", "A/B Testing", "Pandas",
        ],
        "companies": [
            "Deloitte", "Accenture", "KPMG", "EY", "Amazon", "Flipkart",
            "Walmart", "Swiggy", "PwC", "Capgemini",
        ],
        "responsibilities": [
            "Built {count} {bi_tool} dashboards tracking {domain} performance for a team of {team_size} stakeholders.",
            "Wrote {db} queries against a {volume}-row warehouse to size the {domain} opportunity, informing a {pct}% budget shift.",
            "Automated a weekly {domain} report that previously took {hours} hours of manual work.",
            "Ran an A/B test on {domain} conversion that lifted the metric by {pct}% at {confidence}% significance.",
            "Cleaned and reconciled {volume} rows of {domain} data across {count} source systems.",
            "Built a cohort retention model in {tool} that identified a {pct}% drop-off at week {count}.",
            "Presented monthly {domain} findings to {team_size} business leads and translated them into tracked actions.",
            "Defined {count} core metrics and their SQL definitions so reporting stopped disagreeing across teams.",
            "Investigated a {pct}% anomaly in {domain} volumes and traced it to a broken upstream ETL job.",
            "Partnered with engineering to instrument {count} new events for {domain} funnel analysis.",
        ],
        "projects": [
            ("Sales Performance Dashboard", "{bi_tool} dashboard over {volume} transactions with drill-down by region."),
            ("Customer Churn Analysis", "Cohort analysis in {tool} that isolated a {pct}% churn driver."),
            ("Pricing Elasticity Study", "Regression study on {domain} pricing across {count} product lines."),
            ("Marketing Attribution Model", "Multi-touch attribution across {count} channels using {tool}."),
        ],
    },
    "ML Engineer": {
        "titles": [
            "Machine Learning Engineer", "ML Engineer",
            "Applied Scientist", "Senior ML Engineer",
        ],
        "skills": [
            "Python", "Machine Learning", "Deep Learning", "NLP", "PyTorch",
            "TensorFlow", "Scikit-learn", "MLOps", "Docker", "SQL",
        ],
        "companies": [
            "Google", "Microsoft", "Amazon", "NVIDIA", "IBM", "Intel",
            "Adobe", "Samsung R&D", "Qualcomm", "Fractal Analytics",
        ],
        "responsibilities": [
            "Trained and shipped a {model_type} model for {domain}, improving {metric_name} from {pct_low}% to {pct}%.",
            "Built an NLP pipeline over {volume} documents for {domain} classification using {framework}.",
            "Cut inference latency from {latency_before}ms to {latency_after}ms through quantisation and batching.",
            "Set up experiment tracking and model registry so {team_size} engineers could reproduce runs.",
            "Built the feature pipeline feeding {count} features from {db} into the {domain} model.",
            "Ran error analysis on {domain} predictions and closed the top {count} failure modes.",
            "Deployed the {domain} model behind a REST endpoint serving {volume} predictions per day.",
            "Wrote drift monitoring that alerted on a {pct}% distribution shift in {domain} inputs.",
            "Fine-tuned a pretrained {framework} model on {volume} labelled {domain} examples.",
            "Reduced training cost by {pct}% by moving from full retraining to incremental updates.",
        ],
        "projects": [
            ("Resume Classification System", "{framework} text classifier over {volume} documents, {pct}% F1."),
            ("Recommendation Engine", "Collaborative filtering on {volume} interactions serving {domain}."),
            ("Sentiment Analysis Service", "Fine-tuned transformer exposed as a REST endpoint."),
            ("Demand Forecasting Model", "Time series model reducing forecast error by {pct}%."),
        ],
    },
    "Backend Developer": {
        "titles": [
            "Backend Developer", "Backend Engineer",
            "API Engineer", "Senior Backend Developer",
        ],
        "skills": [
            "Python", "FastAPI", "REST APIs", "PostgreSQL", "MongoDB",
            "Docker", "Redis", "Node.js", "Message Queues", "Git",
        ],
        "companies": [
            "Razorpay", "Paytm", "Swiggy", "Zomato", "Flipkart", "Amazon",
            "PhonePe", "CRED", "Meesho", "Zerodha",
        ],
        "responsibilities": [
            "Designed and shipped {count} REST endpoints for the {domain} service in {lang}.",
            "Modelled the {domain} schema in {db} and tuned indexes to cut query time by {pct}%.",
            "Added {cache} caching in front of the {domain} read path, absorbing {volume} requests per day.",
            "Implemented idempotent {domain} processing with a message queue to survive retries.",
            "Built {auth} authentication and rate limiting for {count} public API consumers.",
            "Instrumented the {domain} service with metrics and alerts, cutting incident detection to under {count} minutes.",
            "Wrote integration tests covering {pct}% of the {domain} endpoints and gated deploys on them.",
            "Led the migration of {domain} data from {db} to a partitioned schema without downtime.",
            "Containerised the {domain} stack and moved CI/CD onto {orchestrator}.",
            "Reduced background job failure rate from {pct_low}% to under {pct_tiny}% by adding dead-letter handling.",
        ],
        "projects": [
            ("Payment Processing API", "Idempotent {lang} payment flow over {db} with retry-safe webhooks."),
            ("Food Delivery Backend", "Order and dispatch services handling {volume} orders per day."),
            ("Healthcare Records API", "FHIR-style REST API with {auth} auth and audit logging."),
            ("URL Metrics Service", "High-throughput counter service backed by {cache}."),
        ],
    },
    "Data Scientist": {
        "titles": [
            "Data Scientist", "Applied Data Scientist",
            "Senior Data Scientist", "Quantitative Analyst",
        ],
        "skills": [
            "Python", "Machine Learning", "SQL", "Statistics", "Pandas",
            "Scikit-learn", "Data Visualization", "A/B Testing",
            "Experimentation", "NumPy",
        ],
        "companies": [
            "Amazon", "Microsoft", "Google", "IBM", "Flipkart", "Walmart",
            "Adobe", "Uber", "Ola", "Myntra",
        ],
        "responsibilities": [
            "Built a {model_type} model for {domain} that improved {metric_name} by {pct}% over the existing baseline.",
            "Designed and analysed {count} experiments on {domain}, of which {count_small} shipped.",
            "Performed exploratory analysis on {volume} rows of {domain} data to size a new product bet.",
            "Built causal impact estimates for the {domain} launch when a clean A/B split was not possible.",
            "Productionised the {domain} scoring pipeline with {team_size} engineers and monitored it post-launch.",
            "Segmented {volume} customers into {count} behavioural clusters that drove the {domain} campaign.",
            "Replaced a heuristic {domain} rule with a learned model, cutting false positives by {pct}%.",
            "Wrote the metric definitions and power calculations for the {domain} experimentation framework.",
            "Communicated {domain} model limitations and failure modes to non-technical leadership.",
            "Backtested the {domain} model on {count} historical quarters before rollout.",
        ],
        "projects": [
            ("Fraud Detection Model", "Imbalanced classification over {volume} transactions, {pct}% recall."),
            ("Customer Segmentation", "Clustering {volume} customers into {count} actionable segments."),
            ("Sales Forecasting", "Hierarchical forecast across {count} regions."),
            ("Churn Prediction", "Survival model predicting churn {count} weeks ahead."),
        ],
    },
}


# ============================================================
# SKILL SURFACE FORMS (vocabulary drift)
# ============================================================
# canonical name -> how it might actually be written on a resume.
# The canonical name goes in ground truth and in JDs. The surface form goes in
# the document. Bridging the two is dense retrieval's job.

SKILL_SURFACES: dict[str, list[str]] = {
    "Python": ["Python", "Python 3", "Python (pandas, asyncio)"],
    "Java": ["Java", "Java 17", "Core Java"],
    "SQL": ["SQL", "advanced SQL", "SQL (PostgreSQL, MySQL)"],
    "REST APIs": ["REST APIs", "RESTful services", "REST/HTTP API design"],
    "Docker": ["Docker", "containerisation (Docker)", "Docker & docker-compose"],
    "Kubernetes": ["Kubernetes", "K8s", "container orchestration (Kubernetes)"],
    "PostgreSQL": ["PostgreSQL", "Postgres", "PostgreSQL (partitioning, indexing)"],
    "MongoDB": ["MongoDB", "Mongo", "document databases (MongoDB)"],
    "Machine Learning": ["Machine Learning", "applied ML", "ML modelling"],
    "Deep Learning": ["Deep Learning", "neural networks", "DL"],
    "NLP": ["NLP", "Natural Language Processing", "text mining"],
    "PyTorch": ["PyTorch", "Torch", "PyTorch (training & inference)"],
    "TensorFlow": ["TensorFlow", "TF/Keras", "TensorFlow 2"],
    "Scikit-learn": ["Scikit-learn", "sklearn", "scikit-learn pipelines"],
    "Power BI": ["Power BI", "Microsoft Power BI", "Power BI (DAX)"],
    "Tableau": ["Tableau", "Tableau Desktop", "Tableau dashboards"],
    "Statistics": ["Statistics", "statistical inference", "applied statistics"],
    "Data Visualization": ["Data Visualization", "data viz", "dashboarding"],
    "A/B Testing": ["A/B Testing", "split testing", "online experimentation"],
    "Experimentation": ["Experimentation", "experiment design", "causal inference"],
    "ETL": ["ETL", "data pipelines", "ELT pipelines"],
    "MLOps": ["MLOps", "ML platform engineering", "model deployment & monitoring"],
    "Message Queues": ["Message Queues", "Kafka / RabbitMQ", "async messaging"],
    "Unit Testing": ["Unit Testing", "pytest", "automated testing"],
    "Microservices": ["Microservices", "service-oriented architecture", "microservice design"],
    "Redis": ["Redis", "Redis caching", "in-memory caching (Redis)"],
    "FastAPI": ["FastAPI", "FastAPI + Pydantic", "async Python APIs (FastAPI)"],
    "Node.js": ["Node.js", "NodeJS", "Node (Express)"],
    "Excel": ["Excel", "Advanced Excel", "Excel (pivot tables, Power Query)"],
    "Pandas": ["Pandas", "pandas / numpy", "Pandas dataframes"],
    "NumPy": ["NumPy", "numpy", "NumPy vectorisation"],
    "Git": ["Git", "Git / GitHub", "version control (Git)"],
}


def surface_for(skill: str, rng: random.Random) -> str:
    return rng.choice(SKILL_SURFACES.get(skill, [skill]))


# ============================================================
# TEMPLATE SLOT FILLING
# ============================================================

def slot_values(rng: random.Random) -> dict:
    return {
        "count": rng.randint(3, 24),
        "count_small": rng.randint(1, 5),
        "team_size": rng.randint(2, 12),
        "pct": rng.randint(15, 65),
        "pct_low": rng.randint(4, 30),
        "pct_tiny": rng.choice(["0.5", "1", "2"]),
        "hours": rng.randint(3, 20),
        "confidence": rng.choice([90, 95, 99]),
        "volume": rng.choice([
            "12K", "40K", "150K", "1.2M", "3M", "8M", "220K", "600K",
        ]),
        "latency_before": rng.choice([320, 480, 610, 900, 1400]),
        "latency_after": rng.choice([45, 80, 110, 160, 210]),
        "domain": rng.choice([
            "payments", "logistics", "onboarding", "catalog", "billing",
            "claims", "search", "notifications", "inventory", "risk",
        ]),
        "lang": rng.choice(["Python", "Java", "Go", "Node.js", "Kotlin"]),
        "db": rng.choice(["PostgreSQL", "MySQL", "MongoDB", "Snowflake", "BigQuery"]),
        "cache": rng.choice(["Redis", "Memcached", "an in-process LRU"]),
        "orchestrator": rng.choice(["Kubernetes", "ECS", "GitHub Actions", "Jenkins"]),
        "auth": rng.choice(["OAuth2", "JWT-based", "SAML", "API-key"]),
        "bi_tool": rng.choice(["Power BI", "Tableau", "Looker", "Metabase"]),
        "tool": rng.choice(["Python", "SQL", "R", "Excel"]),
        "framework": rng.choice(["PyTorch", "TensorFlow", "Hugging Face Transformers", "spaCy"]),
        "model_type": rng.choice(["gradient boosting", "logistic regression", "transformer", "random forest"]),
        "metric_name": rng.choice(["precision", "recall", "F1", "AUC", "accuracy"]),
    }


def fill(template: str, rng: random.Random) -> str:
    return template.format(**slot_values(rng))


# ============================================================
# EDUCATION
# ============================================================
# duration in years, so timelines stay coherent.

UG_DEGREES = [
    ("B.Tech in Computer Science", 4),
    ("B.Tech in Information Technology", 4),
    ("B.E. in Computer Engineering", 4),
    ("BCA", 3),
    ("B.Sc in Computer Science", 3),
]

PG_DEGREES = [
    ("M.Tech in Computer Science", 2),
    ("MCA", 2),
    ("MBA in Business Analytics", 2),
    ("M.S. in Data Science", 2),
]

# Prestige tiers exist only as a hidden proxy attribute for the bias audit.
# Tier 3 uses fictional institution names deliberately.
UNIVERSITIES = {
    "tier_1": [
        "IIT Delhi", "IIT Bombay", "IIT Madras", "NIT Trichy",
        "BITS Pilani", "IIIT Hyderabad",
    ],
    "tier_2": [
        "Delhi Technological University", "Anna University",
        "Savitribai Phule Pune University", "VIT Vellore",
        "Manipal Institute of Technology", "Jadavpur University",
    ],
    "tier_3": [
        "Sunrise Institute of Technology", "Greenfield College of Engineering",
        "Northgate Institute of Science", "Lakeview Institute of Technology",
        "Crestwood College of Engineering",
    ],
}

SCHOOL_BOARDS = [
    ("Senior Secondary (Class XII), CBSE", "Delhi Public School"),
    ("Higher Secondary (Class 12), ICSE", "St. Xavier's School"),
    ("Class XII, State Board", "Kendriya Vidyalaya"),
    ("Intermediate (12th)", "DAV Public School"),
]

CERTIFICATIONS = [
    "AWS Certified Solutions Architect - Associate",
    "Google Professional Data Engineer",
    "Microsoft Certified: Azure Data Scientist Associate",
    "Databricks Certified Data Engineer",
    "Certified Kubernetes Application Developer (CKAD)",
    "Tableau Desktop Specialist",
]


# ============================================================
# FORMAT VARIATION
# ============================================================
# These aliases must match the chunker's CANONICAL_SECTIONS map, so that the
# map actually gets exercised. UNKNOWN_SECTIONS deliberately does NOT -- an
# unrecognised heading is a failure mode the chunker has to handle explicitly
# instead of silently appending the content to the previous section.

HEADING_ALIASES = {
    "SUMMARY": ["PROFESSIONAL SUMMARY", "SUMMARY", "PROFILE"],
    "SKILLS": ["SKILLS", "TECHNICAL SKILLS", "CORE COMPETENCIES"],
    "EXPERIENCE": ["PROFESSIONAL EXPERIENCE", "WORK EXPERIENCE", "EXPERIENCE", "EMPLOYMENT HISTORY"],
    "EDUCATION": ["EDUCATION", "ACADEMIC BACKGROUND", "ACADEMICS"],
    "PROJECTS": ["PROJECTS", "PERSONAL PROJECTS", "ACADEMIC PROJECTS"],
    "CERTIFICATIONS": ["CERTIFICATIONS", "CERTIFICATION"],
    "ACHIEVEMENTS": ["ACHIEVEMENTS", "AWARDS", "HONORS & AWARDS"],
}

UNKNOWN_SECTIONS = [
    ("LANGUAGES", ["English (fluent)", "Hindi (native)", "German (basic)"]),
    ("VOLUNTEER EXPERIENCE", ["Mentored students at a local coding bootcamp on weekends."]),
    ("INTERESTS", ["Chess, long-distance running, amateur astronomy."]),
    ("REFERENCES", ["Available on request."]),
]

ACHIEVEMENT_TEMPLATES = [
    "Winner, {domain} track at the national {tool} hackathon.",
    "Awarded internal Engineering Excellence award for the {domain} rewrite.",
    "Ranked in the top {pct_low}% of {count} participants in an inter-college coding contest.",
    "Speaker at a regional meetup on {domain} system design.",
]

DATE_STYLES = ["year_only", "mon_year", "slash", "year_to"]
MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
               "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def format_date_range(start: tuple[int, int], end: tuple[int, int] | None, style: str) -> str:
    """end=None means the candidate is currently in this role."""
    sy, sm = start
    if style == "year_only":
        left, right = str(sy), (str(end[0]) if end else "Present")
        return f"{left} - {right}"
    if style == "mon_year":
        left = f"{MONTH_NAMES[sm - 1]} {sy}"
        right = f"{MONTH_NAMES[end[1] - 1]} {end[0]}" if end else "Present"
        return f"{left} \u2013 {right}"
    if style == "slash":
        left = f"{sm:02d}/{sy}"
        right = f"{end[1]:02d}/{end[0]}" if end else "Present"
        return f"{left} - {right}"
    left, right = str(sy), (str(end[0]) if end else "Present")
    return f"{left} to {right}"


# ============================================================
# CANDIDATE GENERATION
# ============================================================
# Timelines are built BACKWARDS from today so they cannot contradict each
# other: fix the job durations first, anchor the last job to the present, then
# derive graduation year, then derive school years. The old approach assigned
# education years by list index, which produced an M.Tech four years before the
# B.Tech and degrees overlapping with full-time jobs.

NAME_LOCALES = ["en_IN", "en_US", "en_GB", "es_ES", "fr_FR", "de_DE", "pt_BR"]


def generate_candidate(candidate_number: int, rng: random.Random) -> dict:
    role_family = rng.choice(list(CAREER_PROFILES.keys()))
    profile = CAREER_PROFILES[role_family]

    # --- hidden proxy attributes (ground truth only) ---
    locale = rng.choice(NAME_LOCALES)
    local_fake = Faker(locale)
    local_fake.seed_instance(rng.randint(0, 10**9))
    university_tier = rng.choices(
        ["tier_1", "tier_2", "tier_3"], weights=[0.25, 0.45, 0.30]
    )[0]

    # --- jobs: durations and gaps first, positions later ---
    # Everything below is in ABSOLUTE MONTHS (year*12 + month). Working in
    # whole years and then drawing a random month per date lets consecutive
    # jobs overlap by a few months, which is incoherent and would poison any
    # "total years of experience" the extraction stage computes.
    n_jobs = rng.randint(2, 5)
    job_specs = []
    for _ in range(n_jobs):
        job_specs.append({
            "duration_months": rng.randint(12, 48),
            "gap_months": rng.choice([0, 0, 0, 0, 1, 2, 6, 14, 20]),  # ~33% gap
        })
    job_specs[-1]["gap_months"] = 0

    currently_employed = rng.random() < 0.75
    trailing_gap_months = 0 if currently_employed else rng.randint(6, 24)

    span_months = sum(j["duration_months"] + j["gap_months"] for j in job_specs)
    last_end_abs = CURRENT_YEAR * 12 + rng.randint(0, 11) - trailing_gap_months
    career_start_abs = last_end_abs - span_months
    career_start_year = career_start_abs // 12

    # --- education, anchored to career start ---
    has_pg = rng.random() < 0.35
    education = []
    cursor = career_start_year

    if has_pg:
        pg_name, pg_years = rng.choice(PG_DEGREES)
        pg_end, pg_start = cursor, cursor - pg_years
        education.append({
            "degree": pg_name, "level": "postgraduate",
            "institution": rng.choice(UNIVERSITIES[university_tier]),
            "start_year": pg_start, "end_year": pg_end,
        })
        cursor = pg_start

    ug_name, ug_years = rng.choice(UG_DEGREES)
    ug_end, ug_start = cursor, cursor - ug_years
    education.append({
        "degree": ug_name, "level": "undergraduate",
        "institution": rng.choice(UNIVERSITIES[university_tier]),
        "start_year": ug_start, "end_year": ug_end,
    })

    include_school = rng.random() < 0.4
    if include_school:
        school_degree, school_name = rng.choice(SCHOOL_BOARDS)
        education.append({
            "degree": school_degree, "level": "school",
            "institution": school_name,
            "start_year": ug_start - 2, "end_year": ug_start,
        })

    # Real resumes are usually reverse-chronological, but not always.
    reverse_chronological = rng.random() < 0.8
    if not reverse_chronological:
        education = list(reversed(education))

    # --- fill in the job positions on the timeline ---
    companies = rng.sample(profile["companies"], n_jobs)
    date_style = rng.choice(DATE_STYLES)
    experience = []
    cursor_abs = career_start_abs

    def to_ym(abs_months: int) -> tuple[int, int]:
        return abs_months // 12, (abs_months % 12) + 1

    for idx, (spec, company) in enumerate(zip(job_specs, companies)):
        start_abs = cursor_abs
        end_abs = start_abs + spec["duration_months"]
        start_year, start_month = to_ym(start_abs)
        end_year, end_month = to_ym(end_abs)
        is_current = currently_employed and idx == n_jobs - 1

        experience.append({
            "company": company,
            "title": rng.choice(profile["titles"]),
            "start_year": start_year,
            "start_month": start_month,
            "end_year": None if is_current else end_year,
            "end_month": None if is_current else end_month,
            "is_current": is_current,
            "duration_months": spec["duration_months"],
            "date_text": format_date_range(
                (start_year, start_month),
                None if is_current else (end_year, end_month),
                date_style,
            ),
            "responsibilities": [
                fill(t, rng)
                for t in rng.sample(profile["responsibilities"], rng.randint(3, 5))
            ],
        })
        cursor_abs = end_abs + spec["gap_months"]

    if rng.random() < 0.8:
        experience = list(reversed(experience))  # reverse-chronological

    total_experience = sum(j["duration_months"] for j in job_specs) // 12
    has_career_gap = (
        any(j["gap_months"] >= 6 for j in job_specs) or trailing_gap_months >= 6
    )

    # --- skills: canonical for ground truth, surface forms for the document ---
    canonical_skills = rng.sample(profile["skills"], rng.randint(5, 8))
    surface_skills = [surface_for(s, rng) for s in canonical_skills]

    # --- projects ---
    n_projects = rng.randint(1, 3)
    projects = []
    for name, desc_template in rng.sample(profile["projects"], n_projects):
        projects.append({"name": name, "description": fill(desc_template, rng)})

    certifications = rng.sample(CERTIFICATIONS, rng.randint(1, 3))
    achievements = [fill(t, rng) for t in rng.sample(ACHIEVEMENT_TEMPLATES, rng.randint(1, 2))]

    closing_lines = [
        "Enjoys owning a problem end to end, from design through on-call.",
        "Particularly interested in measurable impact rather than shipped tickets.",
        "Works well with product and design on ambiguous problems.",
        "Comfortable mentoring and reviewing across a small team.",
    ]

    summary = (
        f"{experience[0]['title']} with {total_experience}+ years building "
        f"{rng.choice(['production', 'customer-facing', 'internal', 'large-scale'])} systems. "
        f"Comfortable across {', '.join(surface_skills[:3])}. "
        f"{rng.choice(closing_lines)}"
    )

    # --- which optional sections appear, and in what order ---
    sections = ["SUMMARY", "SKILLS", "EXPERIENCE", "EDUCATION"]
    if projects and rng.random() < 0.8:
        sections.append("PROJECTS")
    if rng.random() < 0.6:
        sections.append("CERTIFICATIONS")
    if rng.random() < 0.35:
        sections.append("ACHIEVEMENTS")

    core = sections[:4]
    optional = sections[4:]
    rng.shuffle(optional)
    if rng.random() < 0.3:  # some resumes lead with education
        core = ["SUMMARY", "SKILLS", "EDUCATION", "EXPERIENCE"]
    section_order = core + optional

    unknown_section = None
    if rng.random() < 0.3:
        name, lines = rng.choice(UNKNOWN_SECTIONS)
        unknown_section = {"heading": name, "lines": lines}

    headings = {s: rng.choice(HEADING_ALIASES[s]) for s in HEADING_ALIASES}

    return {
        "id": candidate_number,
        "name": local_fake.name(),
        "email": local_fake.email(),
        "phone": local_fake.phone_number(),
        "city": local_fake.city(),
        "role_family": role_family,
        "current_title": experience[0]["title"],
        "summary": summary,
        "canonical_skills": canonical_skills,
        "surface_skills": surface_skills,
        "education": education,
        "experience": experience,
        "projects": projects,
        "certifications": certifications,
        "achievements": achievements,
        "total_years_experience": total_experience,
        "highest_degree_level": "postgraduate" if has_pg else "undergraduate",
        # hidden attributes for the bias audit -- never printed in the document
        "proxy_attributes": {
            "name_locale": locale,
            "university_tier": university_tier,
            "has_career_gap": has_career_gap,
            "currently_employed": currently_employed,
        },
        # rendering choices, recorded so parser failures can be traced to a cause
        "format": {
            "section_order": section_order,
            "headings": headings,
            "unknown_section": unknown_section,
            "date_style": date_style,
            "reverse_chronological_education": reverse_chronological,
            "include_school": include_school,
            "title_date_layout": rng.choice(["same_line", "separate_lines", "date_then_title"]),
            "bullet_char": None,      # filled in by the renderer
            "heading_style": None,    # filled in by the renderer (docx only)
        },
    }


# ============================================================
# PDF RENDERING
# ============================================================

def create_pdf(candidate: dict, rng: random.Random, fonts: tuple[str, str, list[str]]) -> Path:
    regular, bold, bullets = fonts
    bullet = rng.choice(bullets)
    candidate["format"]["bullet_char"] = bullet
    candidate["format"]["heading_style"] = "bold_larger_text"

    filename = OUTPUT_DIR / f"resume_{candidate['id']:03d}.pdf"
    doc = SimpleDocTemplate(
        str(filename), pagesize=A4,
        rightMargin=45, leftMargin=45, topMargin=40, bottomMargin=40,
    )
    base = getSampleStyleSheet()

    name_style = ParagraphStyle("Name", parent=base["Title"], fontName=bold,
                                alignment=TA_CENTER, fontSize=19, spaceAfter=4)
    role_style = ParagraphStyle("Role", parent=base["Normal"], fontName=regular,
                                alignment=TA_CENTER, fontSize=11, spaceAfter=6)
    contact_style = ParagraphStyle("Contact", parent=base["Normal"], fontName=regular,
                                   alignment=TA_CENTER, fontSize=9, spaceAfter=4)
    head_style = ParagraphStyle("Head", parent=base["Heading2"], fontName=bold,
                                fontSize=12, spaceBefore=11, spaceAfter=4)
    body_style = ParagraphStyle("Body", parent=base["Normal"], fontName=regular,
                                fontSize=9.5, leading=13.5)

    fmt = candidate["format"]
    h = fmt["headings"]
    story: list = []

    story.append(Paragraph(candidate["name"], name_style))
    story.append(Paragraph(candidate["current_title"], role_style))
    story.append(Paragraph(
        f"{candidate['email']} | {candidate['phone']} | {candidate['city']}",
        contact_style,
    ))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(thickness=1))

    def heading(text: str) -> None:
        story.append(Paragraph(text, head_style))

    def render_experience() -> None:
        heading(h["EXPERIENCE"])
        for exp in candidate["experience"]:
            layout = fmt["title_date_layout"]
            if layout == "same_line":
                story.append(Paragraph(
                    f"<b>{exp['title']}</b>, {exp['company']} ({exp['date_text']})",
                    body_style,
                ))
            elif layout == "separate_lines":
                story.append(Paragraph(f"<b>{exp['title']}</b> - {exp['company']}", body_style))
                story.append(Paragraph(exp["date_text"], body_style))
            else:  # date_then_title
                story.append(Paragraph(exp["date_text"], body_style))
                story.append(Paragraph(f"<b>{exp['title']}</b>, {exp['company']}", body_style))
            for r in exp["responsibilities"]:
                story.append(Paragraph(f"{bullet} {r}", body_style))
            story.append(Spacer(1, 5))

    def render_education() -> None:
        heading(h["EDUCATION"])
        for edu in candidate["education"]:
            story.append(Paragraph(
                f"<b>{edu['degree']}</b><br/>{edu['institution']} "
                f"({edu['start_year']} - {edu['end_year']})",
                body_style,
            ))
            story.append(Spacer(1, 4))

    def render_projects() -> None:
        heading(h["PROJECTS"])
        for p in candidate["projects"]:
            story.append(Paragraph(f"<b>{p['name']}</b>", body_style))
            story.append(Paragraph(p["description"], body_style))
            story.append(Spacer(1, 4))

    renderers = {
        "SUMMARY": lambda: (heading(h["SUMMARY"]),
                            story.append(Paragraph(candidate["summary"], body_style))),
        "SKILLS": lambda: (heading(h["SKILLS"]),
                           story.append(Paragraph(
                               f" {bullet} ".join(candidate["surface_skills"]), body_style))),
        "EXPERIENCE": render_experience,
        "EDUCATION": render_education,
        "PROJECTS": render_projects,
        "CERTIFICATIONS": lambda: (heading(h["CERTIFICATIONS"]),
                                   [story.append(Paragraph(f"{bullet} {c}", body_style))
                                    for c in candidate["certifications"]]),
        "ACHIEVEMENTS": lambda: (heading(h["ACHIEVEMENTS"]),
                                 [story.append(Paragraph(f"{bullet} {a}", body_style))
                                  for a in candidate["achievements"]]),
    }

    for section in fmt["section_order"]:
        renderers[section]()

    if fmt["unknown_section"]:
        heading(fmt["unknown_section"]["heading"])
        for line in fmt["unknown_section"]["lines"]:
            story.append(Paragraph(line, body_style))

    doc.build(story)
    return filename


# ============================================================
# DOCX RENDERING
# ============================================================

def create_docx(candidate: dict, rng: random.Random, fonts: tuple[str, str, list[str]]) -> Path:
    bullet = rng.choice(fonts[2])
    candidate["format"]["bullet_char"] = bullet

    # Half of real DOCX resumes never use a real Heading style -- they just bold
    # an uppercase paragraph. A style-based heading detector must survive that.
    heading_style = rng.choice(["real_heading", "bold_caps_paragraph"])
    candidate["format"]["heading_style"] = heading_style

    filename = OUTPUT_DIR / f"resume_{candidate['id']:03d}.docx"
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    sec.left_margin = sec.right_margin = Inches(0.7)

    fmt = candidate["format"]
    h = fmt["headings"]

    def centered(text: str, size: int, bold_text: bool) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold_text
        run.font.size = Pt(size)

    def heading(text: str) -> None:
        if heading_style == "real_heading":
            doc.add_heading(text, level=2)
        else:
            p = doc.add_paragraph()
            run = p.add_run(text.upper())
            run.bold = True
            run.font.size = Pt(12)

    centered(candidate["name"], 19, True)
    centered(candidate["current_title"], 11, False)
    centered(f"{candidate['email']} | {candidate['phone']} | {candidate['city']}", 9, False)

    def render_experience() -> None:
        heading(h["EXPERIENCE"])
        for exp in candidate["experience"]:
            layout = fmt["title_date_layout"]
            p = doc.add_paragraph()
            if layout == "same_line":
                run = p.add_run(f"{exp['title']}, {exp['company']}")
                run.bold = True
                p.add_run(f"  ({exp['date_text']})")
            elif layout == "separate_lines":
                run = p.add_run(f"{exp['title']} - {exp['company']}")
                run.bold = True
                doc.add_paragraph(exp["date_text"])
            else:  # date_then_title
                p.add_run(exp["date_text"])
                p2 = doc.add_paragraph()
                run = p2.add_run(f"{exp['title']}, {exp['company']}")
                run.bold = True
            for r in exp["responsibilities"]:
                doc.add_paragraph(r, style="List Bullet")

    def render_education() -> None:
        heading(h["EDUCATION"])
        for edu in candidate["education"]:
            p = doc.add_paragraph()
            run = p.add_run(edu["degree"])
            run.bold = True
            doc.add_paragraph(
                f"{edu['institution']} ({edu['start_year']} - {edu['end_year']})"
            )

    def render_projects() -> None:
        heading(h["PROJECTS"])
        for proj in candidate["projects"]:
            p = doc.add_paragraph()
            run = p.add_run(proj["name"])
            run.bold = True
            doc.add_paragraph(proj["description"])

    renderers = {
        "SUMMARY": lambda: (heading(h["SUMMARY"]), doc.add_paragraph(candidate["summary"])),
        "SKILLS": lambda: (heading(h["SKILLS"]),
                           doc.add_paragraph(f" {bullet} ".join(candidate["surface_skills"]))),
        "EXPERIENCE": render_experience,
        "EDUCATION": render_education,
        "PROJECTS": render_projects,
        "CERTIFICATIONS": lambda: (heading(h["CERTIFICATIONS"]),
                                   [doc.add_paragraph(c, style="List Bullet")
                                    for c in candidate["certifications"]]),
        "ACHIEVEMENTS": lambda: (heading(h["ACHIEVEMENTS"]),
                                 [doc.add_paragraph(a, style="List Bullet")
                                  for a in candidate["achievements"]]),
    }

    for section in fmt["section_order"]:
        renderers[section]()

    if fmt["unknown_section"]:
        heading(fmt["unknown_section"]["heading"])
        for line in fmt["unknown_section"]["lines"]:
            doc.add_paragraph(line)

    doc.save(filename)
    return filename


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    parser = argparse.ArgumentParser(description="Generate synthetic resumes + ground truth.")
    parser.add_argument("--pdfs", type=int, default=30)
    parser.add_argument("--docxs", type=int, default=20)
    parser.add_argument("--seed", type=int, default=42,
                        help="Fixed seed so the corpus is reproducible across runs.")
    args = parser.parse_args()

    rng = random.Random(args.seed)
    Faker.seed(args.seed)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    GROUND_TRUTH_FILE.parent.mkdir(parents=True, exist_ok=True)

    fonts = register_fonts()
    print(f"Font: {fonts[0]} | usable bullets: {' '.join(fonts[2])}\n")

    candidates = []
    total = args.pdfs + args.docxs

    for i in range(1, total + 1):
        candidate = generate_candidate(i, rng)
        if i <= args.pdfs:
            path = create_pdf(candidate, rng, fonts)
            candidate["file_format"] = "pdf"
        else:
            path = create_docx(candidate, rng, fonts)
            candidate["file_format"] = "docx"
        candidate["file_name"] = path.name
        candidate["resume_id"] = path.stem
        candidates.append(candidate)

        print(
            f"{candidate['file_format'].upper():4} {i:03d} | {candidate['name'][:22]:22} | "
            f"{candidate['role_family']:18} | {candidate['total_years_experience']:2}y | "
            f"jobs={len(candidate['experience'])} | "
            f"dates={candidate['format']['date_style']:9} | "
            f"layout={candidate['format']['title_date_layout']}"
        )

    with open(GROUND_TRUTH_FILE, "w", encoding="utf-8") as f:
        json.dump(candidates, f, indent=2, ensure_ascii=False)

    print(f"\n{total} resumes -> {OUTPUT_DIR}")
    print(f"Ground truth -> {GROUND_TRUTH_FILE}")


if __name__ == "__main__":
    main()